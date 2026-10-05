"""Formal-letter export (GET /api/letter, GET /api/letters/{id}/export): plain text or DOCX with the disclaimer;
placeholders are left as written."""

import io
import re

import docx

from backend.app.models import ParsedResponse

# The prompt asks the model to end with the disclaimer (D14); the API adds its own, so the model's line is removed.
_DISCLAIMER_LINE = re.compile(r"^.*legal information,? not legal advice.*$", re.IGNORECASE | re.MULTILINE)


def without_disclaimer(parsed: ParsedResponse) -> ParsedResponse:
    """Sections with any model-written disclaimer line removed."""
    return parsed.model_copy(
        update={key: _DISCLAIMER_LINE.sub("", getattr(parsed, key)).strip() for key in ("rights", "steps", "letter")}
    )


def text_export(body: str, disclaimer: str) -> str:
    """A letter's text followed by the disclaimer."""
    return f"{body.strip()}\n\n---\n{disclaimer}\n"


def docx_export(body: str, disclaimer: str, title: str = "Formal letter") -> bytes:
    """A .docx with one paragraph per letter line and the disclaimer in the page footer."""
    document = docx.Document()
    document.core_properties.author = "HakiAI"
    document.core_properties.title = title
    for line in body.strip().split("\n"):
        document.add_paragraph(line)
    document.sections[0].footer.paragraphs[0].text = disclaimer
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def letter_text(parsed: ParsedResponse, disclaimer: str) -> str:
    """The letter section followed by the disclaimer."""
    return text_export(parsed.letter, disclaimer)


def letter_docx(parsed: ParsedResponse, disclaimer: str) -> bytes:
    """The letter section as a .docx with the disclaimer in the page footer."""
    return docx_export(parsed.letter, disclaimer)

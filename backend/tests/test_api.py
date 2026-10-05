"""API endpoints via create_app on the fake stack: validation, query, stream, sources, letter, feedback, health."""

import dataclasses
import io
import json
import logging
from pathlib import Path
from typing import Any

import docx
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from backend.app.config import DISCLAIMER, FALLBACK_MESSAGE, Settings
from backend.app.deps import default_deps, real_deps
from backend.app.devstack import fake_deps
from backend.app.generation.llm import OllamaUnavailable
from backend.app.logging_setup import configure_logging
from backend.app.main import create_app
from backend.app.retrieval.meta import IndexMismatchError
from backend.tests.api_support import (
    APP_HEADERS,
    EN_QUESTION,
    GARBAGE,
    SW_QUESTION,
    FakeClock,
    api_settings,
    make_app,
    runtime,
    sse_events,
    sw_translators,
    tokens_text,
)
from backend.tests.fakes import DEFAULT_SCRIPT, FakeLLM

LETTER_SCRIPT = [
    "## RIGHTS EXPLANATION\nAn employer must prove a valid reason (Sample Employment Act, s. 4).\n",
    "## RECOMMENDED STEPS\n1. Write to your employer.\n",
    "## FORMAL LE",
    "TTER\n[Date]\n\nDear [Recipient],\n\nI ask for my unpaid wages.\n\nYours faithfully,\n[Your Name]\n\n",
    DISCLAIMER,
]


def _client(app: FastAPI) -> TestClient:
    return TestClient(app, raise_server_exceptions=False, headers=APP_HEADERS)


def _ask(client: TestClient, question: str = EN_QUESTION, language: str = "en") -> dict[str, Any]:
    response = client.post("/api/query", json={"question": question, "language": language})
    assert response.status_code == 200, response.text
    body: dict[str, Any] = response.json()
    return body


def _stream(client: TestClient, session_id: str) -> list[tuple[str, dict[str, Any]]]:
    response = client.get(f"/api/stream/{session_id}")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("text/event-stream")
    return sse_events(response.text)


def _assert_error(response: Any, status: int, code: str) -> None:
    assert response.status_code == status, response.text
    error = response.json()["error"]
    assert error["code"] == code
    assert error["message"]


@pytest.mark.parametrize(
    "body",
    [
        {"question": "", "language": "en"},
        {"question": "   \n ", "language": "en"},
        {"question": "x" * 1001, "language": "en"},
        {"question": "Why was I fired?", "language": "fr"},
        {"language": "en"},
    ],
)
def test_query_rejects_invalid_bodies_with_error_schema(tmp_path: Path, body: dict[str, Any]) -> None:
    with _client(make_app(tmp_path, FakeLLM())) as client:
        response = client.post("/api/query", json=body)
    _assert_error(response, 422, "invalid_request")
    assert "xxxxxxxxxx" not in response.text


def test_query_rejects_malformed_json_and_control_only_questions(tmp_path: Path) -> None:
    with _client(make_app(tmp_path, FakeLLM())) as client:
        malformed = client.post("/api/query", content=b"{not json", headers={"content-type": "application/json"})
        control = client.post("/api/query", json={"question": "\x00\x01​\x7f", "language": "en"})
    _assert_error(malformed, 422, "invalid_request")
    _assert_error(control, 422, "empty_question")


def test_request_body_limit(tmp_path: Path) -> None:
    with _client(make_app(tmp_path, FakeLLM(), max_body_bytes=200)) as client:
        response = client.post("/api/query", json={"question": "a " * 150, "language": "en"})
    _assert_error(response, 413, "body_too_large")


def test_query_rate_limit_returns_429_with_retry_after(tmp_path: Path) -> None:
    clock = FakeClock()
    with _client(make_app(tmp_path, FakeLLM(), clock=clock, rate_limit_per_min=2)) as client:
        _ask(client)
        _ask(client)
        limited = client.post("/api/query", json={"question": EN_QUESTION, "language": "en"})
        _assert_error(limited, 429, "rate_limited")
        assert limited.headers["retry-after"] == "30"
        clock.now += 30
        _ask(client)


def test_happy_path_streams_tokens_then_done(tmp_path: Path) -> None:
    llm = FakeLLM()
    with _client(make_app(tmp_path, llm)) as client:
        query = _ask(client)
        events = _stream(client, query["session_id"])
    assert query["null_response"] is False
    assert query["language"] == "en"
    assert query["acts"][0] == "sample-employment-act"
    names = [name for name, _ in events]
    assert names[:2] == ["status", "status"]
    assert [data["stage"] for name, data in events if name == "status"] == ["retrieved", "generating"]
    assert names[-1] == "done"
    assert tokens_text(events) == "".join(DEFAULT_SCRIPT)
    deltas = [d for name, data in events if name == "token" for d in data["deltas"]]
    assert {d["section"] for d in deltas} == {"rights", "steps", "letter"}
    done = events[-1][1]
    assert done["format_ok"] is True
    assert done["citation_check"]["unmatched"]
    assert done["warnings"]  # the scripted citation is not in the retrieved chunks
    assert done["disclaimer"] == DISCLAIMER
    assert done["untranslated"] == []
    assert isinstance(done["truncated_chunks"], list)
    assert len(llm.prompts) == 1
    assert EN_QUESTION in llm.prompts[0]
    assert "[CHUNK 1]" in llm.prompts[0]


def test_matched_citations_give_no_warning(tmp_path: Path) -> None:
    with _client(make_app(tmp_path, FakeLLM(LETTER_SCRIPT))) as client:
        events = _stream(client, _ask(client)["session_id"])
    done = events[-1][1]
    assert done["citation_check"] == {"verified": ["Sample Employment Act s. 4"], "unmatched": []}
    assert done["warnings"] == []


def test_sources_match_session_chunks(tmp_path: Path) -> None:
    app = make_app(tmp_path, FakeLLM())
    with _client(app) as client:
        sid = _ask(client)["session_id"]
        sources = client.get(f"/api/sources/{sid}").json()
    session = runtime(app).sessions.get(sid)
    assert session is not None
    assert sources["session_id"] == sid
    assert [c["chunk_id"] for c in sources["chunks"]] == [c.chunk.chunk_id for c in session.chunks]
    assert [c["rank"] for c in sources["chunks"]] == list(range(1, len(session.chunks) + 1))
    first = sources["chunks"][0]
    assert set(first) == {
        "chunk_id",
        "act",
        "unit_type",
        "section_num",
        "section_title",
        "part",
        "page",
        "text",
        "truncated",
        "rank",
    }
    assert first["text"] == session.chunks[0].chunk.text


def test_sources_include_rerank_score_only_in_debug_mode(tmp_path: Path) -> None:
    with _client(make_app(tmp_path, FakeLLM(), debug_scores=True)) as client:
        sid = _ask(client)["session_id"]
        chunks = client.get(f"/api/sources/{sid}").json()["chunks"]
    assert all(isinstance(c["rerank_score"], float) for c in chunks)


def test_letter_txt_and_docx_strip_model_disclaimer_and_keep_placeholders(tmp_path: Path) -> None:
    with _client(make_app(tmp_path, FakeLLM(LETTER_SCRIPT))) as client:
        sid = _ask(client)["session_id"]
        _stream(client, sid)
        txt = client.get(f"/api/letter/{sid}?format=txt")
        default = client.get(f"/api/letter/{sid}")
        docx_reply = client.get(f"/api/letter/{sid}?format=docx")
    assert txt.status_code == 200
    assert txt.headers["content-type"].startswith("text/plain")
    assert "attachment" in txt.headers["content-disposition"]
    assert default.text == txt.text
    assert txt.text.startswith("[Date]\n\nDear [Recipient],")
    assert "[Your Name]" in txt.text
    assert txt.text.count(DISCLAIMER) == 1
    assert txt.text.rstrip().endswith(DISCLAIMER)
    assert docx_reply.headers["content-type"].startswith("application/vnd.openxmlformats")
    document = docx.Document(io.BytesIO(docx_reply.content))
    body = "\n".join(p.text for p in document.paragraphs)
    assert "Dear [Recipient]," in body
    assert "[Your Name]" in body
    assert DISCLAIMER not in body
    assert document.sections[0].footer.paragraphs[0].text == DISCLAIMER


def test_unknown_sessions_are_404(tmp_path: Path) -> None:
    with _client(make_app(tmp_path, FakeLLM())) as client:
        for path in ("/api/stream/nope", "/api/sources/nope", "/api/letter/nope"):
            _assert_error(client.get(path), 404, "session_not_found")
        _assert_error(
            client.post("/api/feedback", json={"session_id": "nope", "rating": "up"}), 404, "session_not_found"
        )


def test_letter_conflicts_and_bad_format(tmp_path: Path) -> None:
    with _client(make_app(tmp_path, FakeLLM())) as client:
        sid = _ask(client)["session_id"]
        _assert_error(client.get(f"/api/letter/{sid}"), 409, "not_finished")
        _assert_error(client.get(f"/api/letter/{sid}?format=pdf"), 422, "invalid_request")
        null_sid = _ask(client, GARBAGE)["session_id"]
        _assert_error(client.get(f"/api/letter/{null_sid}"), 404, "no_letter")


def test_letter_missing_section_is_404(tmp_path: Path) -> None:
    with _client(make_app(tmp_path, FakeLLM(["## RIGHTS EXPLANATION\nOnly rights.\n"]))) as client:
        sid = _ask(client)["session_id"]
        events = _stream(client, sid)
        assert events[-1][1]["format_ok"] is False
        _assert_error(client.get(f"/api/letter/{sid}"), 404, "no_letter")


def test_second_connection_replays_without_calling_llm_again(tmp_path: Path) -> None:
    llm = FakeLLM()
    with _client(make_app(tmp_path, llm)) as client:
        sid = _ask(client)["session_id"]
        first = _stream(client, sid)
        second = _stream(client, sid)
    assert len(llm.prompts) == 1
    assert tokens_text(second) == tokens_text(first)
    assert second[-1] == first[-1]
    replay_deltas = {d["section"]: d["text"] for d in second[0][1]["deltas"]}
    assert replay_deltas["steps"] == "1. Write to the other party."


def test_stream_while_running_is_409(tmp_path: Path) -> None:
    app = make_app(tmp_path, FakeLLM())
    with _client(app) as client:
        sid = _ask(client)["session_id"]
        session = runtime(app).sessions.get(sid)
        assert session is not None
        session.status = "running"
        _assert_error(client.get(f"/api/stream/{sid}"), 409, "in_progress")
        session.status = "aborted"
        _assert_error(client.get(f"/api/stream/{sid}"), 409, "aborted")


def test_stream_is_503_when_the_queue_is_full(tmp_path: Path) -> None:
    app = make_app(tmp_path, FakeLLM(), max_queue=0)
    with _client(app) as client:
        sid = _ask(client)["session_id"]
        holder = runtime(app).gate.enter()
        busy = client.get(f"/api/stream/{sid}")
        _assert_error(busy, 503, "busy")
        assert busy.headers["retry-after"]
        holder.release()
        assert _stream(client, sid)[-1][0] == "done"


def test_null_path_never_invokes_llm(tmp_path: Path) -> None:
    llm = FakeLLM()
    with _client(make_app(tmp_path, llm)) as client:
        query = _ask(client, GARBAGE)
        events = _stream(client, query["session_id"])
        again = _stream(client, query["session_id"])
        sources = client.get(f"/api/sources/{query['session_id']}").json()
    assert query["null_response"] is True
    assert events == again == [("null", {"message": FALLBACK_MESSAGE, "disclaimer": DISCLAIMER})]
    assert sources["chunks"] == []
    assert llm.prompts == []


def test_llm_failure_emits_error_event_and_blocks_replay(tmp_path: Path) -> None:
    llm = FakeLLM(fail_with=OllamaUnavailable("Ollama is not reachable at http://localhost:11434"))
    with _client(make_app(tmp_path, llm)) as client:
        sid = _ask(client)["session_id"]
        events = _stream(client, sid)
        assert events[-1] == (
            "error",
            {"code": "llm_unavailable", "message": "Ollama is not reachable at http://localhost:11434"},
        )
        _assert_error(client.get(f"/api/stream/{sid}"), 409, "generation_failed")


def test_swahili_flow_emits_translated_sections_with_citations_intact(tmp_path: Path) -> None:
    llm = FakeLLM(LETTER_SCRIPT)
    app = make_app(tmp_path, llm, **sw_translators())
    with _client(app) as client:
        query = _ask(client, SW_QUESTION, "sw")
        events = _stream(client, query["session_id"])
        letter = client.get(f"/api/letter/{query['session_id']}").text
    assert query["language"] == "sw"
    assert "[en]" in llm.prompts[0]
    stages = [data["stage"] for name, data in events if name == "status"]
    assert stages == ["retrieved", "generating", "translating"]
    names = [name for name, _ in events]
    assert names.index("translated") == len(names) - 2
    sections = dict(events)["translated"]["sections"]
    assert "[sw]" in sections["rights"]
    assert "Sample Employment Act, s. 4" in sections["rights"]
    assert "[Your Name]" in sections["letter"]
    assert DISCLAIMER not in sections["letter"]
    done = events[-1][1]
    ui_sw = json.loads(Settings().ui_strings_path.read_text(encoding="utf-8"))["sw"]
    assert done["disclaimer"] == ui_sw["disclaimer"]
    assert ui_sw["translation_note"] in done["warnings"]
    assert "[sw]" in letter
    assert letter.rstrip().endswith(ui_sw["disclaimer"])


def test_swahili_null_message_is_in_swahili(tmp_path: Path) -> None:
    app = make_app(tmp_path, FakeLLM(), **sw_translators())
    with _client(app) as client:
        events = _stream(client, _ask(client, GARBAGE, "sw")["session_id"])
    ui_sw = json.loads(Settings().ui_strings_path.read_text(encoding="utf-8"))["sw"]
    assert events == [("null", {"message": ui_sw["fallback"], "disclaimer": ui_sw["disclaimer"]})]


def test_feedback_appends_once_per_session_without_content(tmp_path: Path) -> None:
    llm = FakeLLM()
    with _client(make_app(tmp_path, llm)) as client:
        sid = _ask(client)["session_id"]
        _stream(client, sid)
        first = client.post("/api/feedback", json={"session_id": sid, "rating": "down", "comment": "Too\x00 vague"})
        second = client.post("/api/feedback", json={"session_id": sid, "rating": "up"})
        too_long = client.post("/api/feedback", json={"session_id": sid, "rating": "up", "comment": "x" * 501})
        bad_rating = client.post("/api/feedback", json={"session_id": sid, "rating": "meh"})
    assert first.json() == {"recorded": True}
    assert second.json() == {"recorded": False}
    _assert_error(too_long, 422, "invalid_request")
    _assert_error(bad_rating, 422, "invalid_request")
    lines = (tmp_path / "feedback.jsonl").read_text(encoding="utf-8").splitlines()
    assert len(lines) == 1
    record = json.loads(lines[0])
    assert set(record) == {"timestamp", "session_id", "rating", "comment", "language", "null_response", "chunk_ids"}
    assert (record["session_id"], record["rating"], record["comment"]) == (sid, "down", "Too vague")
    assert (record["language"], record["null_response"]) == ("en", False)
    assert EN_QUESTION not in lines[0]
    assert "Write to the other party" not in lines[0]


def test_feedback_rate_limited(tmp_path: Path) -> None:
    with _client(make_app(tmp_path, FakeLLM(), clock=FakeClock(), rate_limit_per_min=1)) as client:
        sid = _ask(client)["session_id"]
        client.post("/api/feedback", json={"session_id": sid, "rating": "up"})
        _assert_error(client.post("/api/feedback", json={"session_id": sid, "rating": "up"}), 429, "rate_limited")


def test_health_ok(tmp_path: Path) -> None:
    with _client(make_app(tmp_path, FakeLLM())) as client:
        response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json() == {
        "status": "ok",
        "ollama": True,
        "model_present": True,
        "indexes_loaded": True,
        "models_warm": True,
    }


@pytest.mark.parametrize(
    ("llm", "failed"),
    [(FakeLLM(healthy=False), {"ollama", "model_present"}), (FakeLLM(model_present=False), {"model_present"})],
)
def test_health_503_when_ollama_or_model_missing(tmp_path: Path, llm: FakeLLM, failed: set[str]) -> None:
    with _client(make_app(tmp_path, llm)) as client:
        response = client.get("/api/health")
    body = response.json()
    assert response.status_code == 503
    assert body["status"] == "degraded"
    assert {key for key in ("ollama", "model_present", "indexes_loaded", "models_warm") if not body[key]} == failed
    assert body["error"]["code"] == "degraded"
    assert "ollama" in body["error"]["message"]


def test_health_and_query_503_before_startup(tmp_path: Path) -> None:
    client = _client(make_app(tmp_path, FakeLLM()))  # no `with`: lifespan never ran
    health = client.get("/api/health").json()
    assert (health["indexes_loaded"], health["models_warm"]) == (False, False)
    _assert_error(client.post("/api/query", json={"question": EN_QUESTION}), 503, "not_ready")


def test_startup_fails_fast_with_rebuild_hint(tmp_path: Path) -> None:
    deps = fake_deps(api_settings(tmp_path), tmp_path / "work", llm=FakeLLM())

    def broken() -> Any:
        raise IndexMismatchError("no readable dense index; build it with: python -m backend.app.ingestion.build_index")

    app = create_app(dataclasses.replace(deps, load_pipeline=broken))
    with pytest.raises(IndexMismatchError, match="build_index"), TestClient(app):
        pass


def test_startup_wraps_missing_files_with_actionable_message(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    for var in ("HF_HUB_OFFLINE", "TRANSFORMERS_OFFLINE"):
        monkeypatch.setenv(var, "1")  # restored after the test; the real loader only setdefaults them
    deps = real_deps(api_settings(tmp_path, chunks_path=tmp_path / "missing.json"), llm=FakeLLM())
    with pytest.raises(RuntimeError, match=r"build_index.*setup_offline"), TestClient(create_app(deps)):
        pass


def test_session_expiry(tmp_path: Path) -> None:
    clock = FakeClock()
    with _client(make_app(tmp_path, FakeLLM(), clock=clock, session_ttl_s=60)) as client:
        sid = _ask(client)["session_id"]
        assert client.get(f"/api/sources/{sid}").status_code == 200
        clock.now += 61
        _assert_error(client.get(f"/api/sources/{sid}"), 404, "session_not_found")


def test_no_content_in_logs(tmp_path: Path) -> None:
    buffer = io.StringIO()
    root = logging.getLogger()
    saved = (root.handlers[:], root.level)
    configure_logging(Settings(), buffer)
    try:
        app = make_app(tmp_path, FakeLLM(LETTER_SCRIPT), **sw_translators())
        with _client(app) as client:
            for question, lang in ((EN_QUESTION, "en"), (SW_QUESTION, "sw"), (GARBAGE, "en")):
                sid = _ask(client, question, lang)["session_id"]
                _stream(client, sid)
                client.post("/api/feedback", json={"session_id": sid, "rating": "up", "comment": "secret note"})
    finally:
        root.handlers[:], level = saved
        root.setLevel(level)
    logged = buffer.getvalue()
    assert "session_id" in logged
    assert "generation_done" in logged
    for content in (EN_QUESTION, SW_QUESTION, GARBAGE, "unpaid wages", "Dear [Recipient]", "secret note", "employer"):
        assert content not in logged


def test_security_headers_and_no_cors_by_default(tmp_path: Path) -> None:
    with _client(make_app(tmp_path, FakeLLM())) as client:
        response = client.get("/api/health", headers={"Origin": "http://localhost:5173"})
    assert response.headers["x-content-type-options"] == "nosniff"
    assert response.headers["x-frame-options"] == "DENY"
    assert response.headers["cache-control"] == "no-store"
    assert "default-src 'self'" in response.headers["content-security-policy"]
    assert "access-control-allow-origin" not in response.headers


def test_cors_only_for_dev_origin_in_dev_mode(tmp_path: Path) -> None:
    with _client(make_app(tmp_path, FakeLLM(), dev_mode=True)) as client:
        allowed = client.get("/api/health", headers={"Origin": "http://localhost:5173"})
        other = client.get("/api/health", headers={"Origin": "http://evil.example"})
    assert allowed.headers["access-control-allow-origin"] == "http://localhost:5173"
    assert "access-control-allow-origin" not in other.headers


def test_static_frontend_with_spa_fallback_after_api_routes(tmp_path: Path) -> None:
    dist = tmp_path / "dist"
    (dist / "assets").mkdir(parents=True)
    (dist / "index.html").write_text("<html>HakiAI</html>", encoding="utf-8")
    (dist / "assets" / "app.js").write_text("console.log(1)", encoding="utf-8")
    with _client(make_app(tmp_path, FakeLLM(), frontend_dist=dist)) as client:
        assert client.get("/").text == "<html>HakiAI</html>"
        assert client.get("/some/client/route").text == "<html>HakiAI</html>"
        assert client.get("/assets/app.js").text == "console.log(1)"
        assert client.get("/api/health").json()["status"] == "ok"
        _assert_error(client.get("/api/unknown"), 404, "not_found")
        docs = client.get("/docs")
        assert docs.status_code == 200
        assert "content-security-policy" not in docs.headers
        assert "/api/query" in client.get("/openapi.json").json()["paths"]


def test_fake_backends_flag_selects_the_fake_stack(tmp_path: Path) -> None:
    deps = default_deps(api_settings(tmp_path, fake_backends=True))
    assert deps.settings.acts[0].name == "Sample Constitution"
    assert deps.setup_logging is True
    with _client(create_app(dataclasses.replace(deps, setup_logging=False))) as client:
        assert client.get("/api/health").status_code == 200
        assert _ask(client)["null_response"] is False


def test_real_deps_use_configured_corpus_without_loading_models(tmp_path: Path) -> None:
    deps = default_deps(api_settings(tmp_path))
    assert deps.settings.acts == Settings().acts
    assert deps.setup_logging is True

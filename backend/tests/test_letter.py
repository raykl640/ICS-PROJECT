"""Letter export and removal of the model-appended disclaimer line."""

import io

import docx

from backend.app.config import DISCLAIMER
from backend.app.letter import letter_docx, letter_text, without_disclaimer
from backend.app.models import ParsedResponse

PARSED = ParsedResponse(
    rights="You have rights.",
    steps="1. Write.",
    letter="[Date]\n\nDear [Recipient],\n\nI write about my wages.\n\nYours faithfully,\n[Your Name]",
)


def test_without_disclaimer_drops_the_line_in_any_section() -> None:
    parsed = ParsedResponse(
        rights=f"You have rights.\n{DISCLAIMER}",
        steps="1. Write.",
        letter=f"Dear [Recipient],\n\n**{DISCLAIMER}**\n",
        format_ok=False,
    )
    cleaned = without_disclaimer(parsed)
    assert cleaned == ParsedResponse(
        rights="You have rights.", steps="1. Write.", letter="Dear [Recipient],", format_ok=False
    )


def test_letter_text_appends_the_given_disclaimer() -> None:
    text = letter_text(PARSED, "Taarifa, si ushauri.")
    assert text.startswith(PARSED.letter)
    assert text.endswith("Taarifa, si ushauri.\n")


def test_letter_docx_keeps_lines_and_placeholders_with_disclaimer_footer() -> None:
    document = docx.Document(io.BytesIO(letter_docx(PARSED, DISCLAIMER)))
    lines = [p.text for p in document.paragraphs]
    assert lines == PARSED.letter.split("\n")
    assert document.sections[0].footer.paragraphs[0].text == DISCLAIMER
    assert document.core_properties.author == "HakiAI"

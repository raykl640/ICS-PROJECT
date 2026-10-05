"""Library endpoints: conversations, letters (versions, restore, export), matters, bookmarks, notes; search, filters,
cursor pages; cross-user isolation; nothing readable at rest."""

import io
from pathlib import Path
from typing import Any

import docx
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from backend.app.config import DISCLAIMER
from backend.tests.api_support import APP_HEADERS, EN_QUESTION, TickingClock, make_app
from backend.tests.fakes import FakeLLM

PW = "a long enough password"
MARKER = "Nyokabi-Plaintext-Marker-5521"
CHUNK = "sample-employment-act-4"
LETTER = "Dear Sir or Madam,"


def client(app: FastAPI) -> TestClient:
    return TestClient(app, raise_server_exceptions=False, headers=APP_HEADERS)


def register(c: TestClient, username: str) -> TestClient:
    r = c.post("/api/auth/register", json={"username": username, "password": PW})
    assert r.status_code == 200, r.text
    return c


@pytest.fixture
def app(tmp_path: Path) -> FastAPI:
    return make_app(
        tmp_path,
        FakeLLM(),
        wall_clock=TickingClock(),
        app_db_path=tmp_path / "app.db",
        letter_versions_max=3,
        library_page_size=2,
    )


@pytest.fixture
def amina(app: FastAPI) -> Any:
    with client(app) as c:  # runs the app's lifespan once; other clients share the running app
        yield register(c, "amina")


def ask_and_finish(c: TestClient, question: str = EN_QUESTION, **extra: Any) -> dict[str, Any]:
    """A saved turn: ask without background and read the stream to the end (the turn is saved after it)."""
    body: dict[str, Any] = c.post("/api/query", json={"question": question, "background": False, **extra}).json()
    assert c.get(f"/api/stream/{body['session_id']}").status_code == 200
    detail: dict[str, Any] = c.get(f"/api/conversations/{body['conversation_id']}").json()
    return detail


def code(r: Any) -> str:
    value: str = r.json()["error"]["code"]
    return value


def test_conversation_rename_pin_file_and_delete(amina: TestClient) -> None:
    conv = ask_and_finish(amina)
    matter = amina.post("/api/matters", json={"name": "Dismissal at the bakery"}).json()
    r = amina.patch(f"/api/conversations/{conv['id']}", json={"title": "  My  dismissal ", "pinned": True})
    assert (r.json()["title"], r.json()["pinned"]) == ("My dismissal", True)
    r = amina.patch(f"/api/conversations/{conv['id']}", json={"matter_id": matter["id"]})
    assert r.json()["matter_id"] == matter["id"] and r.json()["title"] == "My dismissal"
    assert amina.patch(f"/api/conversations/{conv['id']}", json={"matter_id": None}).json()["matter_id"] is None
    assert code(amina.patch(f"/api/conversations/{conv['id']}", json={"title": "  "})) == "empty_title"
    assert amina.delete(f"/api/conversations/{conv['id']}").json() == {"deleted": True}
    assert amina.get(f"/api/conversations/{conv['id']}").status_code == 404


def test_long_first_question_becomes_an_80_character_title(amina: TestClient) -> None:
    question = "My employer " + "really " * 30 + "fired me without notice."
    conv = ask_and_finish(amina, question)
    assert len(conv["title"]) == 80 and conv["title"].endswith("…")
    assert conv["thread"][0]["question"] == question


def test_conversation_search_filters_sort_and_cursor_pages(amina: TestClient) -> None:
    first = ask_and_finish(amina, "My employer fired me without notice and kept my wages")
    second = ask_and_finish(amina, "My landlord locked me out of the shop")
    third = ask_and_finish(amina, "Wages were not paid by my employer this month")
    amina.patch(f"/api/conversations/{second['id']}", json={"pinned": True})

    page1 = amina.get("/api/conversations").json()
    assert [c["id"] for c in page1["items"]] == [second["id"], third["id"]]  # pinned first, then newest
    page2 = amina.get("/api/conversations", params={"cursor": page1["next_cursor"]}).json()
    assert [c["id"] for c in page2["items"]] == [first["id"]] and page2["next_cursor"] is None

    found = amina.get("/api/conversations", params={"q": "WAGES"}).json()["items"]
    assert {c["id"] for c in found} == {first["id"], third["id"]}
    assert amina.get("/api/conversations", params={"pinned": "true"}).json()["items"][0]["id"] == second["id"]
    assert amina.get("/api/conversations", params={"lang": "sw"}).json()["items"] == []
    act = first["acts"][0]
    assert all(act in c["acts"] for c in amina.get("/api/conversations", params={"act": act}).json()["items"])
    titles = amina.get("/api/conversations", params={"sort": "title", "pinned": "false"}).json()["items"]
    assert [c["title"] for c in titles] == sorted((first["title"], third["title"]), key=str.casefold)
    assert amina.get("/api/conversations", params={"from": "2999-01-01"}).json()["items"] == []
    assert code(amina.get("/api/conversations", params={"cursor": "%%%"})) == "invalid_cursor"


def test_letter_from_an_answer_keeps_versions_up_to_the_cap_and_restores(amina: TestClient) -> None:
    conv = ask_and_finish(amina)
    turn_id = conv["thread"][0]["id"]
    letter = amina.post("/api/letters", json={"title": "To my employer", "turn_id": turn_id}).json()
    assert letter["body"] == LETTER and letter["version"] == 1 and letter["conversation_id"] == conv["id"]
    for n in range(2, 6):
        saved = amina.put(f"/api/letters/{letter['id']}", json={"body": f"{LETTER}\nDraft {n}"}).json()
        assert saved["version"] == n
    assert amina.put(f"/api/letters/{letter['id']}", json={"body": f"{LETTER}\nDraft 5"}).json()["version"] == 5
    versions = amina.get(f"/api/letters/{letter['id']}/versions").json()
    assert [v["n"] for v in versions] == [5, 4, 3]  # letter_versions_max = 3
    restored = amina.post(f"/api/letters/{letter['id']}/versions/3/restore").json()
    assert (restored["version"], restored["body"]) == (6, f"{LETTER}\nDraft 3")
    assert amina.post(f"/api/letters/{letter['id']}/versions/1/restore").status_code == 404
    renamed = amina.put(f"/api/letters/{letter['id']}", json={"title": "Final", "pinned": True}).json()
    assert (renamed["title"], renamed["pinned"], renamed["version"]) == ("Final", True, 6)


def test_blank_letter_limits_and_line_breaks(amina: TestClient, app: FastAPI) -> None:
    blank = amina.post("/api/letters", json={}).json()
    assert (blank["title"], blank["body"], blank["version"]) == ("Letter", "", 1)
    kept = amina.put(f"/api/letters/{blank['id']}", json={"body": "Line one\r\n\tLine two​"}).json()
    assert kept["body"] == "Line one\n Line two"
    too_long = "x" * (app.state.runtime.deps.settings.letter_max_chars + 1)
    assert code(amina.post("/api/letters", json={"body": too_long})) in ("letter_too_long", "body_too_large")
    assert amina.post("/api/letters", json={"turn_id": "nope"}).status_code == 404


def test_letter_export_txt_and_docx_keep_the_disclaimer(amina: TestClient) -> None:
    body = "[Date]\n\nDear [Recipient],\n\nI ask for my wages.\n\nYours faithfully,\n[Your Name]"
    letter = amina.post("/api/letters", json={"title": "Wages", "body": body}).json()
    txt = amina.get(f"/api/letters/{letter['id']}/export", params={"format": "txt"})
    assert txt.headers["content-disposition"] == 'attachment; filename="letter.txt"'
    assert txt.text == f"{body}\n\n---\n{DISCLAIMER}\n"
    sw = amina.get(f"/api/letters/{letter['id']}/export", params={"format": "txt", "lang": "sw"})
    assert DISCLAIMER not in sw.text and sw.text.startswith(body)
    docx_reply = amina.get(f"/api/letters/{letter['id']}/export", params={"format": "docx"})
    document = docx.Document(io.BytesIO(docx_reply.content))
    assert [p.text for p in document.paragraphs] == body.split("\n")
    assert document.sections[0].footer.paragraphs[0].text == DISCLAIMER
    assert document.core_properties.title == "Wages"


def test_matters_collect_items_and_survive_as_folders(amina: TestClient) -> None:
    conv = ask_and_finish(amina)
    matter = amina.post("/api/matters", json={"name": "  Bakery   case "}).json()
    assert (matter["name"], matter["status"]) == ("Bakery case", "open")
    mid = matter["id"]
    amina.patch(f"/api/conversations/{conv['id']}", json={"matter_id": mid})
    letter = amina.post("/api/letters", json={"title": "L", "matter_id": mid}).json()
    bookmark = amina.post("/api/bookmarks", json={"chunk_id": CHUNK, "matter_id": mid}).json()
    note = amina.post("/api/notes", json={"body": "Call the union", "matter_id": mid}).json()
    detail = amina.get(f"/api/matters/{mid}").json()
    assert [c["id"] for c in detail["conversations"]] == [conv["id"]]
    assert [x["id"] for x in detail["letters"]] == [letter["id"]]
    assert [x["id"] for x in detail["bookmarks"]] == [bookmark["id"]]
    assert [x["id"] for x in detail["notes"]] == [note["id"]]
    closed = amina.patch(f"/api/matters/{mid}", json={"status": "closed", "name": "Bakery (done)"}).json()
    assert (closed["status"], closed["name"]) == ("closed", "Bakery (done)")
    other = amina.post("/api/matters", json={"name": "Rent"}).json()
    assert [m["id"] for m in amina.get("/api/matters").json()["items"]] == [other["id"], mid]  # open first
    assert code(amina.post("/api/matters", json={"name": " "})) == "empty_name"
    assert amina.delete(f"/api/matters/{mid}").json() == {"deleted": True}
    assert amina.get(f"/api/conversations/{conv['id']}").json()["matter_id"] is None
    assert amina.get(f"/api/letters/{letter['id']}").json()["matter_id"] is None
    unfiled = amina.get("/api/notes", params={"matter": "none"}).json()["items"]
    assert [n["id"] for n in unfiled] == [note["id"]]


def test_bookmarks_resolve_sections_and_do_not_duplicate(amina: TestClient) -> None:
    first = amina.post("/api/bookmarks", json={"chunk_id": CHUNK}).json()
    assert first["act"] and first["section_num"] == "4" and first["section_title"]
    assert amina.post("/api/bookmarks", json={"chunk_id": CHUNK}).json()["id"] == first["id"]
    assert code(amina.post("/api/bookmarks", json={"chunk_id": "no-such-section-1"})) == "section_not_found"
    matter = amina.post("/api/matters", json={"name": "Work"}).json()
    assert amina.patch(f"/api/bookmarks/{first['id']}", json={"matter_id": matter["id"]}).json()["matter_id"]
    found = amina.get("/api/bookmarks", params={"act": "sample-employment-act"}).json()["items"]
    assert [b["id"] for b in found] == [first["id"]]
    assert amina.get("/api/bookmarks", params={"act": "sample-tenancy-act"}).json()["items"] == []
    assert amina.delete(f"/api/bookmarks/{first['id']}").json() == {"deleted": True}


def test_notes_targets_edit_and_search(amina: TestClient) -> None:
    conv = ask_and_finish(amina)
    note = amina.post(
        "/api/notes", json={"body": "Ask HR\nfor the letter", "target_kind": "conversation", "target_id": conv["id"]}
    ).json()
    assert (note["target_kind"], note["target_id"], note["body"]) == (
        "conversation",
        conv["id"],
        "Ask HR\nfor the letter",
    )
    edited = amina.put(f"/api/notes/{note['id']}", json={"body": "Ask HR today"}).json()
    assert edited["body"] == "Ask HR today"
    assert (
        amina.post("/api/notes", json={"body": "On s.4", "target_kind": "chunk", "target_id": CHUNK}).status_code == 200
    )
    assert code(amina.post("/api/notes", json={"body": "x", "target_kind": "letter"})) == "invalid_target"
    assert code(amina.post("/api/notes", json={"body": "x", "target_id": "abc"})) == "invalid_target"
    assert code(amina.post("/api/notes", json={"body": "  "})) == "empty_note"
    assert [n["body"] for n in amina.get("/api/notes", params={"q": "hr TODAY"}).json()["items"]] == ["Ask HR today"]


def test_another_user_cannot_read_change_or_delete_any_id(app: FastAPI, amina: TestClient) -> None:
    conv = ask_and_finish(amina)
    letter = amina.post("/api/letters", json={"title": "Mine", "body": "text"}).json()
    matter = amina.post("/api/matters", json={"name": "Mine"}).json()
    bookmark = amina.post("/api/bookmarks", json={"chunk_id": CHUNK}).json()
    note = amina.post("/api/notes", json={"body": "Mine"}).json()
    baraka = register(client(app), "baraka")
    attempts = [
        ("GET", f"/api/conversations/{conv['id']}", None),
        ("PATCH", f"/api/conversations/{conv['id']}", {"title": "Stolen"}),
        ("DELETE", f"/api/conversations/{conv['id']}", None),
        ("GET", f"/api/letters/{letter['id']}", None),
        ("PUT", f"/api/letters/{letter['id']}", {"body": "Stolen"}),
        ("GET", f"/api/letters/{letter['id']}/versions", None),
        ("POST", f"/api/letters/{letter['id']}/versions/1/restore", None),
        ("GET", f"/api/letters/{letter['id']}/export", None),
        ("DELETE", f"/api/letters/{letter['id']}", None),
        ("GET", f"/api/matters/{matter['id']}", None),
        ("PATCH", f"/api/matters/{matter['id']}", {"name": "Stolen"}),
        ("DELETE", f"/api/matters/{matter['id']}", None),
        ("PATCH", f"/api/bookmarks/{bookmark['id']}", {"matter_id": None}),
        ("DELETE", f"/api/bookmarks/{bookmark['id']}", None),
        ("PUT", f"/api/notes/{note['id']}", {"body": "Stolen"}),
        ("DELETE", f"/api/notes/{note['id']}", None),
        ("POST", "/api/letters", {"turn_id": conv["thread"][0]["id"]}),
        ("POST", "/api/notes", {"body": "x", "target_kind": "letter", "target_id": letter["id"]}),
        ("POST", "/api/bookmarks", {"chunk_id": CHUNK, "matter_id": matter["id"]}),
    ]
    for method, path, body in attempts:
        r = baraka.request(method, path, json=body)
        assert r.status_code == 404, (method, path, r.text)
    mine = baraka.post("/api/letters", json={"title": "B"}).json()
    r = baraka.put(f"/api/letters/{mine['id']}", json={"matter_id": matter["id"]})
    assert (r.status_code, code(r)) == (404, "matter_not_found")
    for kind in ("conversations", "letters", "matters", "bookmarks", "notes"):
        items = baraka.get(f"/api/{kind}").json()["items"]
        assert {item["id"] for item in items}.isdisjoint(
            {conv["id"], letter["id"], matter["id"], bookmark["id"], note["id"]}
        )
    assert amina.get(f"/api/letters/{letter['id']}").json()["body"] == "text"
    assert amina.get(f"/api/conversations/{conv['id']}").json()["title"] == EN_QUESTION


def test_library_needs_an_unlocked_sign_in(app: FastAPI, amina: TestClient) -> None:
    guest = client(app)
    assert code(guest.get("/api/conversations")) == "auth_required"
    assert code(guest.post("/api/letters", json={})) == "auth_required"
    amina.post("/api/auth/lock")
    assert code(amina.get("/api/letters")) == "locked"


def test_nothing_readable_at_rest_and_export_has_it_all(app: FastAPI, amina: TestClient, tmp_path: Path) -> None:
    conv = ask_and_finish(amina, f"{MARKER} my employer fired me")
    amina.patch(f"/api/conversations/{conv['id']}", json={"title": f"{MARKER} title"})
    letter = amina.post("/api/letters", json={"title": f"{MARKER} letter", "body": f"{MARKER} body"}).json()
    amina.put(f"/api/letters/{letter['id']}", json={"body": f"{MARKER} body v2"})
    matter = amina.post("/api/matters", json={"name": f"{MARKER} matter"}).json()
    amina.post("/api/notes", json={"body": f"{MARKER} note", "target_kind": "matter", "target_id": matter["id"]})
    amina.post("/api/bookmarks", json={"chunk_id": CHUNK})
    stored = b"".join(p.read_bytes() for p in tmp_path.glob("app.db*"))
    for secret in (MARKER, "Write to the other party", CHUNK, LETTER):
        assert secret.encode() not in stored, secret
    exported = amina.get("/api/account/export").json()["library"]
    assert exported["conversations"][0]["thread"][0]["question"].startswith(MARKER)
    assert "sources" not in exported["conversations"][0]["thread"][0]
    assert [v["n"] for v in exported["letters"][0]["versions"]] == [2, 1]
    assert exported["matters"][0]["name"] == f"{MARKER} matter"
    assert exported["notes"][0]["body"] == f"{MARKER} note"
    assert exported["bookmarks"][0]["chunk_id"] == CHUNK


def test_deleting_the_account_deletes_the_library(app: FastAPI, amina: TestClient, tmp_path: Path) -> None:
    ask_and_finish(amina)
    amina.post("/api/letters", json={"body": "text"})
    assert amina.request("DELETE", "/api/account", json={"password": PW}).status_code == 200
    rt = app.state.runtime
    with rt.library.db.read() as conn:
        for table in ("conversations", "turns", "letters", "letter_versions", "notes", "bookmarks", "matters"):
            assert conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0] == 0
